"""Reads locations from the `PSGC` sheet of the PSA publication datafile."""

from collections.abc import Iterator
from dataclasses import dataclass
from functools import cache
from pathlib import Path

from openpyxl import load_workbook

from app.core.config import get_settings
from app.domain.entities.barangay import Barangay, UrbanRural
from app.domain.entities.city import City, CityClass
from app.domain.entities.municipality import Municipality, MunicipalityType
from app.domain.entities.province import Province
from app.domain.entities.region import Region

SHEET_NAME = "PSGC"

PSGC_CODE_COLUMN = "10-digit PSGC"
NAME_COLUMN = "Name"
CORRESPONDENCE_CODE_COLUMN = "Correspondence Code"
LEVEL_COLUMN = "Geographic Level"
OLD_NAMES_COLUMN = "Old names"
CITY_CLASS_COLUMN = "City Class"
INCOME_CLASSIFICATION_COLUMN = "Income\nClassification (DOF DO No. 074.2024)"
URBAN_RURAL_COLUMN = "Urban / Rural\n(based on 2020 CPH)"
POPULATION_COLUMN = "2024 Population"
STATUS_COLUMN = "Status"

REGION_LEVEL = "Reg"
PROVINCE_LEVEL = "Prov"
CITY_LEVEL = "City"
BARANGAY_LEVEL = "Bgy"
MUNICIPALITY_LEVELS = {
    "Mun": MunicipalityType.MUNICIPALITY,
    "SubMun": MunicipalityType.SUB_MUNICIPALITY,
}

CAPITAL_STATUS = "Capital"
NOT_APPLICABLE = "-"

REQUIRED_COLUMNS = frozenset(
    {
        PSGC_CODE_COLUMN,
        NAME_COLUMN,
        CORRESPONDENCE_CODE_COLUMN,
        LEVEL_COLUMN,
        OLD_NAMES_COLUMN,
        CITY_CLASS_COLUMN,
        INCOME_CLASSIFICATION_COLUMN,
        URBAN_RURAL_COLUMN,
        POPULATION_COLUMN,
        STATUS_COLUMN,
    }
)


@dataclass(frozen=True, slots=True)
class PsgcDataset:
    regions: tuple[Region, ...]
    provinces: tuple[Province, ...]
    cities: tuple[City, ...]
    municipalities: tuple[Municipality, ...]
    barangays: tuple[Barangay, ...]


def _clean_code(value: object, *, width: int) -> str | None:
    """Excel stores some codes as numbers, which drops their leading zeros."""
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    return text.zfill(width)


def _clean_str(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text if text and text != NOT_APPLICABLE else None


def _clean_int(value: object) -> int | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    return int(value)


def _iter_rows(path: Path) -> Iterator[dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook[SHEET_NAME].iter_rows(values_only=True)
        header = [str(cell).strip() if cell is not None else "" for cell in next(rows)]
        missing = REQUIRED_COLUMNS.difference(header)
        if missing:
            raise ValueError(f"{path} sheet '{SHEET_NAME}' is missing columns: {sorted(missing)}")
        for row in rows:
            yield dict(zip(header, row, strict=False))
    finally:
        workbook.close()


def _region_code_of(psgc_code: str) -> str:
    """The first two digits of a 10-digit PSGC code identify its region."""
    return psgc_code[:2] + "0" * 8


def _parent_code_of(psgc_code: str) -> str:
    """The first five digits identify the province, or the city for sub-municipalities."""
    return psgc_code[:5] + "0" * 5


def _lgu_code_of(psgc_code: str) -> str:
    """The first seven digits identify a barangay's city, municipality, or sub-municipality."""
    return psgc_code[:7] + "0" * 3


def _to_region(psgc_code: str, row: dict[str, object]) -> Region:
    return Region(
        psgc_code=psgc_code,
        name=str(row[NAME_COLUMN]).strip(),
        correspondence_code=_clean_code(row.get(CORRESPONDENCE_CODE_COLUMN), width=9),
        population=_clean_int(row.get(POPULATION_COLUMN)),
    )


def _to_province(psgc_code: str, row: dict[str, object]) -> Province:
    return Province(
        psgc_code=psgc_code,
        name=str(row[NAME_COLUMN]).strip(),
        correspondence_code=_clean_code(row.get(CORRESPONDENCE_CODE_COLUMN), width=9),
        region_code=_region_code_of(psgc_code),
        income_classification=_clean_str(row.get(INCOME_CLASSIFICATION_COLUMN)),
        population=_clean_int(row.get(POPULATION_COLUMN)),
    )


def _old_names(row: dict[str, object]) -> tuple[str, ...]:
    old_name = _clean_str(row.get(OLD_NAMES_COLUMN))
    return (old_name,) if old_name else ()


def _to_city(psgc_code: str, row: dict[str, object], *, province_codes: set[str]) -> City:
    parent_code = _parent_code_of(psgc_code)
    city_class = _clean_str(row.get(CITY_CLASS_COLUMN))
    try:
        parsed_class = CityClass(city_class)
    except ValueError:
        raise ValueError(f"City {psgc_code} has unknown city class {city_class!r}") from None
    return City(
        psgc_code=psgc_code,
        name=str(row[NAME_COLUMN]).strip(),
        correspondence_code=_clean_code(row.get(CORRESPONDENCE_CODE_COLUMN), width=9),
        city_class=parsed_class,
        region_code=_region_code_of(psgc_code),
        province_code=parent_code if parent_code in province_codes else None,
        income_classification=_clean_str(row.get(INCOME_CLASSIFICATION_COLUMN)),
        is_capital=_clean_str(row.get(STATUS_COLUMN)) == CAPITAL_STATUS,
        old_names=_old_names(row),
        population=_clean_int(row.get(POPULATION_COLUMN)),
    )


def _to_municipality(
    psgc_code: str,
    row: dict[str, object],
    municipality_type: MunicipalityType,
    *,
    province_codes: set[str],
) -> Municipality:
    parent_code = _parent_code_of(psgc_code)
    is_sub_municipality = municipality_type is MunicipalityType.SUB_MUNICIPALITY
    return Municipality(
        psgc_code=psgc_code,
        name=str(row[NAME_COLUMN]).strip(),
        correspondence_code=_clean_code(row.get(CORRESPONDENCE_CODE_COLUMN), width=9),
        type=municipality_type,
        region_code=_region_code_of(psgc_code),
        province_code=parent_code if parent_code in province_codes else None,
        city_code=parent_code if is_sub_municipality else None,
        income_classification=_clean_str(row.get(INCOME_CLASSIFICATION_COLUMN)),
        is_capital=_clean_str(row.get(STATUS_COLUMN)) == CAPITAL_STATUS,
        old_names=_old_names(row),
        population=_clean_int(row.get(POPULATION_COLUMN)),
    )


def _to_barangay(
    psgc_code: str,
    row: dict[str, object],
    *,
    province_codes: set[str],
    city_codes: set[str],
    municipality_types: dict[str, MunicipalityType],
) -> Barangay:
    lgu_code = _lgu_code_of(psgc_code)
    parent_code = _parent_code_of(psgc_code)
    municipality_type = municipality_types.get(lgu_code)
    if lgu_code in city_codes:
        city_code = lgu_code
    elif municipality_type is MunicipalityType.SUB_MUNICIPALITY:
        city_code = parent_code
    else:
        city_code = None
    urban_rural = _clean_str(row.get(URBAN_RURAL_COLUMN))
    return Barangay(
        psgc_code=psgc_code,
        name=str(row[NAME_COLUMN]).strip(),
        correspondence_code=_clean_code(row.get(CORRESPONDENCE_CODE_COLUMN), width=9),
        region_code=_region_code_of(psgc_code),
        province_code=parent_code if parent_code in province_codes else None,
        city_code=city_code,
        municipality_code=lgu_code if municipality_type is not None else None,
        urban_rural=UrbanRural(urban_rural) if urban_rural else None,
        old_names=_old_names(row),
        population=_clean_int(row.get(POPULATION_COLUMN)),
    )


def _check_parents(
    path: Path,
    regions: list[Region],
    provinces: list[Province],
    cities: list[City],
    municipalities: list[Municipality],
    barangays: list[Barangay],
) -> None:
    region_codes = {region.psgc_code for region in regions}
    city_codes = {city.psgc_code for city in cities}
    problems = [
        f"province {p.psgc_code} has no region" for p in provinces if p.region_code not in region_codes
    ]
    problems += [f"city {c.psgc_code} has no region" for c in cities if c.region_code not in region_codes]
    problems += [
        f"municipality {m.psgc_code} has no region"
        for m in municipalities
        if m.region_code not in region_codes
    ]
    problems += [
        f"sub-municipality {m.psgc_code} has no parent city"
        for m in municipalities
        if m.type is MunicipalityType.SUB_MUNICIPALITY and m.city_code not in city_codes
    ]
    problems += [
        f"barangay {b.psgc_code} has no region" for b in barangays if b.region_code not in region_codes
    ]
    problems += [
        f"barangay {b.psgc_code} has no city or municipality"
        for b in barangays
        if b.city_code is None and b.municipality_code is None
    ]
    if problems:
        raise ValueError(f"Inconsistent hierarchy in {path}: {problems}")


def read_dataset(path: Path) -> PsgcDataset:
    regions: list[Region] = []
    provinces: list[Province] = []
    city_rows: list[tuple[str, dict[str, object]]] = []
    municipality_rows: list[tuple[str, dict[str, object], MunicipalityType]] = []
    barangay_rows: list[tuple[str, dict[str, object]]] = []

    for row in _iter_rows(path):
        psgc_code = _clean_code(row.get(PSGC_CODE_COLUMN), width=10)
        if psgc_code is None:
            continue
        level = row.get(LEVEL_COLUMN)
        if level == REGION_LEVEL:
            regions.append(_to_region(psgc_code, row))
        elif level == PROVINCE_LEVEL:
            provinces.append(_to_province(psgc_code, row))
        elif level == CITY_LEVEL:
            city_rows.append((psgc_code, row))
        elif isinstance(level, str) and level in MUNICIPALITY_LEVELS:
            municipality_rows.append((psgc_code, row, MUNICIPALITY_LEVELS[level]))
        elif level == BARANGAY_LEVEL:
            barangay_rows.append((psgc_code, row))

    province_codes = {province.psgc_code for province in provinces}
    cities = [_to_city(psgc_code, row, province_codes=province_codes) for psgc_code, row in city_rows]
    municipalities = [
        _to_municipality(psgc_code, row, municipality_type, province_codes=province_codes)
        for psgc_code, row, municipality_type in municipality_rows
    ]
    city_codes = {city.psgc_code for city in cities}
    municipality_types = {m.psgc_code: m.type for m in municipalities}
    barangays = [
        _to_barangay(
            psgc_code,
            row,
            province_codes=province_codes,
            city_codes=city_codes,
            municipality_types=municipality_types,
        )
        for psgc_code, row in barangay_rows
    ]

    _check_parents(path, regions, provinces, cities, municipalities, barangays)

    return PsgcDataset(
        regions=tuple(regions),
        provinces=tuple(provinces),
        cities=tuple(cities),
        municipalities=tuple(municipalities),
        barangays=tuple(barangays),
    )


@cache
def get_dataset() -> PsgcDataset:
    """Parse the configured datafile once per process; later calls reuse the result."""
    return read_dataset(get_settings().psgc_datafile)
